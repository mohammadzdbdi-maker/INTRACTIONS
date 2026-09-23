#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
تبدیل همه PDFهای مکمل/گیاهی سامانه پیام داروساز به اکسل.

خروجی به‌ازای هر PDF یک فایل .xlsx با این ساختار:
  ردیف ۱ : تاریخ اعلام (همان که بالای PDF نوشته شده، مثلاً «30 شهریور 1405»)
  ردیف ۲ : سرستون‌ها → ردیف | کد پیام | نام فارسی | نام انگلیسی | قیمت
  ردیف ۳+: داده‌ها

ستون «قیمت» = قیمت واحد (ریال). قیمت بسته در PDF وجود دارد ولی طبق
مشخصات درخواستی در خروجی نیامده است.

استخراج: pdfplumber (اصلی) + PyMuPDF (fallback برای PDFهای دارای ساختار خراب)
"""
import os
import re
import sys
import glob
import contextlib
import unicodedata

import pdfplumber
import pymupdf
import pypdf

try:
    pymupdf.TOOLS.mupdf_display_errors(False)
except Exception:
    pass
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

SRC_DIR = "Suppliment PDF"
OUT_DIR = "Suppliment Excel"

FA = lambda s: any('\u0600' <= c <= '\u06FF' or '\uFB50' <= c <= '\uFEFF' for c in s)


def norm(s):
    return (unicodedata.normalize('NFKC', s)
            .replace('\ufeff', '').replace('\u200f', '').replace('\u200e', '')
            .replace('\x00', '').strip())


def fa_w(w):
    """کلمه فارسی: ترتیب کاراکترها در PDF معکوسِ ترتیب منطقی است."""
    return norm(w[::-1])


def tok(w):
    w = re.sub(r'(?<=\d)[\u0600-\u06FF\uFB50-\uFEFF](?=\d)', '', w)
    return fa_w(w) if FA(w) else norm(w)


# ---------------------------------------------------------------- استخراج کلمات
def words_pdfplumber(path):
    pages = []
    with pdfplumber.open(path) as pdf:
        for p in pdf.pages:
            pages.append([(w['x0'], w['top'], w['text']) for w in p.extract_words()])
    return pages


def words_pymupdf(path):
    """fallback: کاراکترها را بر اساس فاصله افقی به کلمه تبدیل می‌کنیم."""
    pages = []
    with pymupdf.open(path) as doc:
        for page in doc:
            ws = page.get_text("words")          # هر گلیف یک «word»
            lines = {}
            for x0, y0, x1, y1, t, *_ in ws:
                lines.setdefault(round(y0 / 3), []).append((x0, y0, x1, y1, t))
            out = []
            for k in sorted(lines):
                cur = None
                for x0, y0, x1, y1, t in sorted(lines[k]):
                    if cur and (x0 - cur[2]) < 1.6 and abs(y0 - cur[1]) < 3:
                        cur = (cur[0], cur[1], x1, y1, cur[4] + t)
                    else:
                        if cur:
                            out.append((cur[0], cur[1], cur[4]))
                        cur = (x0, y0, x1, y1, t)
                if cur:
                    out.append((cur[0], cur[1], cur[4]))
            pages.append(out)
    return pages


# ---------------------------------------------------------------- تقویم
import datetime
import jdatetime

JALALI_MONTHS = ['', 'فروردین', 'اردیبهشت', 'خرداد', 'تیر', 'مرداد', 'شهریور',
                 'مهر', 'آبان', 'آذر', 'دی', 'بهمن', 'اسفند']


def jalali_str_from_filename(base):
    m = re.search(r'(\d{4})(\d{2})?(\d{2})?\.pdf$', base)
    if not m:
        return None
    y, mo, d = m.group(1), m.group(2), m.group(3)
    if not mo:
        return f"{int(y)}"
    if not d:
        return f"{JALALI_MONTHS[int(mo)]} {int(y)}" if mo and 1 <= int(mo) <= 12 else f"{int(y)}"
    jd = jdatetime.date.fromgregorian(date=datetime.date(int(y), int(mo), int(d)))
    return f"{jd.day} {JALALI_MONTHS[jd.month]} {jd.year}"


# ---------------------------------------------------------------- پارس رکوردها
def parse_pages(pages):
    date_str = None
    parts = {'day': None, 'mon': None, 'yr': None}
    recs = []
    for pi, W in enumerate(pages):
        if pi == 0:
            day = [t for x, y, t in W if y < 35 and re.fullmatch(r'\d{1,2}', t)]
            mon = [t for x, y, t in W if 35 <= y < 65 and fa_w(t) in JALALI_MONTHS]
            yr = [t for x, y, t in W if 65 <= y < 100 and re.fullmatch(r'1[34]\d{2}', t)]
            parts = {'day': day[0] if day else None,
                     'mon': fa_w(mon[0]) if mon else None,
                     'yr': yr[0] if yr else None}
            if day and mon and yr:
                date_str = f"{day[0]} {fa_w(mon[0])} {yr[0]}"
        codes = sorted([(y, m.group(1)) for x, y, t in W if 440 < x < 545
                        and (m := re.search(r'(0\d{9})', t))])
        for i, (cy, code) in enumerate(codes):
            prev = codes[i - 1][0] if i > 0 else 150
            nxt = codes[i + 1][0] if i + 1 < len(codes) else 10 ** 6
            lo = (prev + cy) / 2 if i > 0 else 150
            hi = (cy + nxt) / 2 if i + 1 < len(codes) else 10 ** 6
            rw = [(x, y, t) for x, y, t in W if lo <= y < hi]
            rad = [t for x, y, t in rw if x > 545 and re.fullmatch(r'\d+', t)]
            plines = {}
            for x, y, t in rw:
                if x < 160 and y < cy and re.fullmatch(r'\d[\d,]*', t):
                    plines.setdefault(round(y / 3), []).append((x, y, t))
            cand = [v for v in plines.values() if any(x < 100 for x, _, _ in v)] \
                or list(plines.values())
            if cand:
                # خط قیمت = خطی که نزدیک‌ترین y به خط کد دارد
                line = max(cand, key=lambda v: v[0][1])
                b = [t for x, _, t in line if x < 100]
                v = [t for x, _, t in line if 100 <= x < 160]
                if v and b:
                    bast, vahd = b[0], v[0]
                elif b:
                    bast, vahd = None, b[0]      # قالب تک‌قیمتی / قیمت واحد جاافتاده
                elif v:
                    bast, vahd = None, v[0]
                else:
                    bast = vahd = None
            else:
                bast = vahd = None
            fa_t = sorted([(y, x, t) for x, y, t in rw if 200 < x < 470 and y < cy],
                          key=lambda z: (round(z[0] / 3), -z[1]))
            en_t = sorted([(y, x, t) for x, y, t in rw if 100 < x < 470 and y > cy and not FA(t)],
                          key=lambda z: (round(z[0] / 3), z[1]))
            recs.append({
                'radif': int(rad[0]) if rad else None,
                'code': code,
                'fa': ' '.join(tok(t) for _, _, t in fa_t),
                'en': ' '.join(norm(t) for _, _, t in en_t),
                'vahd': vahd,
                'bast': bast,
            })
    return date_str, recs, parts


# ------------------------------------------------------- موتور سوم: pypdf خط‌به‌خط
# برای PDFهایی که dict فونتشان خراب است و pdfminer/PyMuPDF هیچ متن قابل‌استفاده‌ای
# نمی‌دهند؛ pypdf متن را خط‌به‌خط (به‌ترتیب منطقی) برمی‌گرداند.
def _is_fa_line(s):
    return bool(re.search(r'[\u0600-\u06FF\uFB50-\uFEFF]', s)) and not re.search(r'[A-Za-z]{2}', s)


def parse_pdf_pypdf(path):
    r = pypdf.PdfReader(path)
    recs, date_parts, buf = [], [], []

    def build(lines, code, radif):
        # آخرین سه‌گانه «روز/ماه/سال» در بافر = شروع رکورد جاری
        for j in range(len(lines) - 2):
            if (re.fullmatch(r'\d{1,2}', lines[j]) and norm(lines[j + 1]) in JALALI_MONTHS
                    and re.fullmatch(r'1[34]\d{2}', lines[j + 2])):
                if not date_parts:
                    date_parts.extend(lines[j:j + 3])
                break
        # لنگر قیمت: خط «ریال»؛ قیمت عددِ قبل از آن و نام‌ها بعد از آن هستند
        ri, suffix = None, ''
        for k, l in enumerate(lines):
            nl = norm(l)
            if nl in ('ریال', 'لایر', 'رﻳﺎل'):
                ri = k
                break
            if nl.startswith(('ریال', 'لایر', 'رﻳﺎل')) and len(nl) > 4:
                ri, suffix = k, nl[4:].strip()   # «ریال» چسبیده به نام دارو
                break
        if ri is not None:
            price = lines[ri - 1] if ri and re.fullmatch(r'[\d,]+', lines[ri - 1]) else None
            rest = lines[ri + 1:]
            if suffix:
                mf = re.match(r'([\u0600-\u06FF\uFB50-\uFEFF\s]+)(.*)$', suffix)
                if mf and mf.group(2):
                    rest = [mf.group(1).strip(), mf.group(2).strip()] + rest
                else:
                    rest = [suffix] + rest
        else:
            price, rest = None, lines
        fa = ' '.join(l for l in rest if _is_fa_line(l))
        en = ' '.join(l for l in rest if re.search(r'[A-Za-z]', l))
        return {'radif': radif, 'code': code, 'fa': fa, 'en': en, 'vahd': price, 'bast': None}

    for page in r.pages:
        raw = [norm(l) for l in (page.extract_text() or '').splitlines()]
        lines = []
        i = 0
        while i < len(raw):
            if (re.fullmatch(r'0\d{9}', raw[i]) and i + 1 < len(raw)
                    and re.fullmatch(r'\d{1,3}', raw[i + 1])):
                lines.append(raw[i] + ' ' + raw[i + 1]); i += 2
            else:
                lines.append(raw[i]); i += 1
        for line in lines:
            if not line:
                continue
            m = re.search(r'(0\d{9})\s+(\d{1,3})$', line)
            if m:
                prefix = line[:m.start()].strip()   # کد گاهی به انتهای نام انگلیسی چسبیده
                if prefix:
                    buf.append(prefix)
                rec = build(buf, m.group(1), int(m.group(2)))
                if rec['vahd'] and rec['fa']:
                    recs.append(rec)
                buf = []
            else:
                buf.append(line)
    date_str = (' '.join(date_parts[:3]) if len(date_parts) == 3 else None)
    return date_str, recs, 'pypdf'


@contextlib.contextmanager
def quiet_stderr():
    """خطاهای MuPDF/pdfminer روی فایل‌های خراب را پنهان می‌کند."""
    sys.stderr.flush()
    with open(os.devnull, 'w') as dn:
        fd, saved = sys.stderr.fileno(), os.dup(2)
        os.dup2(dn.fileno(), fd)
        try:
            yield
        finally:
            os.dup2(saved, fd)
            os.close(saved)


def parse_pdf(path):
    try:
        pages = words_pdfplumber(path)
        date_str, recs, parts = parse_pages(pages)
        if recs:
            return date_str, recs, 'pdfplumber', parts
    except Exception:
        pass
    try:
        with quiet_stderr():
            pages = words_pymupdf(path)
        date_str, recs, parts = parse_pages(pages)
        if recs:
            return date_str, recs, 'pymupdf', parts
    except Exception:
        pass
    try:
        with quiet_stderr():
            d, r_, e = parse_pdf_pypdf(path)
            return d, r_, e, {'day': None, 'mon': None, 'yr': None}
    except Exception:
        return None, [], 'failed', {'day': None, 'mon': None, 'yr': None}


# ---------------------------------------------------------------- نوشتن اکسل
HDR = ['ردیف', 'کد پیام', 'نام فارسی', 'نام انگلیسی', 'قیمت']


def write_xlsx(path_out, date_str, recs):
    wb = Workbook()
    ws = wb.active
    ws.title = "Supplements"
    ws.sheet_view.rightToLeft = True

    ws['A1'] = date_str or ''
    ws['A1'].font = Font(bold=True, size=12)
    ws['A1'].alignment = Alignment(horizontal='right')

    for c, h in enumerate(HDR, 1):
        cell = ws.cell(row=2, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="009900")
        cell.alignment = Alignment(horizontal='center')

    r = 3
    for rec in recs:
        price = None
        if rec['vahd']:
            try:
                price = int(rec['vahd'].replace(',', ''))
            except ValueError:
                price = rec['vahd']
        ws.cell(row=r, column=1, value=rec['radif'])
        ws.cell(row=r, column=2, value=rec['code'])
        ws.cell(row=r, column=3, value=rec['fa'])
        ws.cell(row=r, column=4, value=rec['en'])
        ws.cell(row=r, column=5, value=price)
        r += 1

    for col, w in zip("ABCDE", [8, 14, 55, 45, 14]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A3"
    wb.save(path_out)


def main():
    only = sys.argv[1] if len(sys.argv) > 1 else None
    os.makedirs(OUT_DIR, exist_ok=True)
    files = sorted(glob.glob(os.path.join(SRC_DIR, "*.pdf")))
    if only:
        files = [f for f in files if only in f]
    total = 0
    problems = []
    manifest = []
    for f in files:
        base = os.path.basename(f)
        date_str, recs, engine, parts = parse_pdf(f)
        if not date_str:
            fj = jalali_str_from_filename(base)
            if fj:
                date_str = fj + "  (از نام فایل)"
                fparts = fj.split()
                # تطبیق ماه/سال هدر با نام فایل (روز در هدر این فایلها نیست)
                if parts['mon'] and parts['yr']:
                    if not (parts['mon'] in fparts and parts['yr'] in fparts):
                        problems.append((base, f"date mismatch: hdr={parts['mon']} {parts['yr']} vs file={fj}", engine))
        if not recs:
            problems.append((base, 'no-records', engine))
        for rec in recs:
            if rec['radif'] is None or not rec['vahd']:
                problems.append((base, f"incomplete row {rec['code']}", engine))
                break
        out = os.path.join(OUT_DIR, base.replace('.pdf', '.xlsx'))
        write_xlsx(out, date_str, recs)
        total += len(recs)
        manifest.append((base, len(recs), engine, date_str or ''))
        print(f"{len(recs):4d}  {engine:10s}  {date_str or '??':<28}  {base}")
    with open(os.path.join('analysis', 'supplement_excel_index.csv'), 'w',
              newline='', encoding='utf-8-sig') as fh:
        import csv as _csv
        w = _csv.writer(fh)
        w.writerow(['pdf', 'xlsx', 'jalali_date', 'records', 'engine'])
        for base, n, eng, d in manifest:
            w.writerow([base, base.replace('.pdf', '.xlsx'), d, n, eng])
    print(f"\nجمع رکورد: {total}  |  فایل: {len(files)}  |  مشکل‌دار: {len(problems)}")
    for p in problems[:30]:
        print("  ⚠️ ", p)


if __name__ == '__main__':
    main()
