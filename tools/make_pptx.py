# -*- coding: utf-8 -*-
"""ساخت نسخه PowerPoint (pptx) قابل ویرایش — اسلایدهای A4 افقی، راست‌به‌چپ، فونت Arad."""
import math
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from content import COVER, PAGES, FOOTER, SEVERITY

OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                   'بروشور-تداخلات-دارویی.pptx')
FONT = 'Arad'
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
DARK = RGBColor(0x2B, 0x2B, 0x2B)
GRAY = RGBColor(0x5C, 0x5C, 0x5C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEAD_BG = RGBColor(0xE3, 0xE3, 0xE3)
ZEBRA = RGBColor(0xF5, 0xF5, 0xF5)
BOX_BG = RGBColor(0xF4, 0xF4, 0xF4)
LINE = RGBColor(0xAD, 0xAD, 0xAD)

SLIDE_W, SLIDE_H = 11.69, 8.27           # A4 افقی بر حسب اینچ
M = 0.45
TOP = 1.02
CONTENT_W = SLIDE_W - 2 * M


def rtl(par, align=PP_ALIGN.RIGHT):
    par.alignment = align
    pPr = par._p.get_or_add_pPr()
    pPr.set('rtl', '1')
    return par


def write(par, text, size=11, bold=False, color=BLACK, font=FONT, space_after=2):
    run = par.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font
    run.font.color.rgb = color
    par.space_after = Pt(space_after)
    return run


def textbox(slide, x, y, w, h, text, size=11, bold=False, color=BLACK,
            align=PP_ALIGN.RIGHT, wrap=True, spacing=1.0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    rtl(p, align)
    p.line_spacing = spacing
    write(p, text, size, bold, color)
    return tb


def roundbox(slide, x, y, w, h, fill=BOX_BG, line=LINE, line_w=0.75):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_w)
    sh.shadow.inherit = False
    sh.text_frame.word_wrap = True
    return sh


def chars_per_line(width_in, size):
    return max(6, int(width_in * 72 / (size * 0.58)))


def est_lines(text, width_in, size):
    cpl = chars_per_line(width_in, size)
    n = 0
    for line in str(text).split('\n'):
        n += max(1, math.ceil(len(line) / cpl))
    return n


def text_h(text, width_in, size, leading=1.32):
    return est_lines(text, width_in, size) * size * leading / 72.0


def add_table(slide, x, y, width, header, rows, widths, fs=9.5):
    ncol = len(header)
    tot = sum(widths) if widths else ncol
    colw = [(width * (w / tot)) for w in (widths or [1] * ncol)]
    # برآورد ارتفاع هر سطر
    heights = [max(1, est_lines(header[i], colw[i] - 0.14, fs)) for i in range(ncol)]
    h_header = max(heights) * fs * 1.55 / 72.0 + 0.12
    row_h = []
    for row in rows:
        lines = max(est_lines(row[i], colw[i] - 0.14, fs) for i in range(ncol))
        row_h.append(lines * fs * 1.5 / 72.0 + 0.12)
    total_h = h_header + sum(row_h)
    gt = slide.shapes.add_table(len(rows) + 1, ncol, Inches(x), Inches(y),
                                Inches(width), Inches(total_h))
    table = gt.table
    for i, cw in enumerate(colw):
        table.columns[i].width = Emu(int(Inches(cw)))
    table.rows[0].height = Emu(int(Inches(h_header)))
    for i, rh in enumerate(row_h):
        table.rows[i + 1].height = Emu(int(Inches(rh)))
    for c in range(ncol):
        cell = table.cell(0, c)
        cell.text = ''
        cell.fill.solid()
        cell.fill.fore_color.rgb = HEAD_BG
        cell.margin_left = cell.margin_right = Inches(0.05)
        cell.margin_top = cell.margin_bottom = Inches(0.02)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        rtl(p)
        write(p, str(header[c]), fs, True, BLACK)
    for r, row in enumerate(rows):
        for c in range(ncol):
            cell = table.cell(r + 1, c)
            cell.text = ''
            cell.fill.solid()
            cell.fill.fore_color.rgb = ZEBRA if r % 2 else RGBColor(0xFF, 0xFF, 0xFF)
            cell.margin_left = cell.margin_right = Inches(0.05)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            rtl(p)
            bold = (c == 0 and str(row[c]) in ('X', 'D', 'C', 'B'))
            write(p, str(row[c]), fs, bold, BLACK)
    return total_h


def note_box(slide, x, y, w, title, body, fs=10.5):
    h_body = text_h(body, w - 0.5, fs, 1.45)
    total = 0.34 + h_body + 0.18
    sh = roundbox(slide, x, y, w, total, BOX_BG, DARK, 1.0)
    tf = sh.text_frame
    tf.margin_right = tf.margin_left = Inches(0.14)
    tf.margin_top = tf.margin_bottom = Inches(0.06)
    p = tf.paragraphs[0]
    rtl(p)
    p.space_after = Pt(3)
    write(p, title, fs + 1, True, BLACK)
    for line in str(body).split('\n'):
        pb = tf.add_paragraph()
        rtl(pb)
        pb.space_after = Pt(2)
        write(pb, line, fs, False, BLACK)
    return total


def card_boxes(slide, x, y, w, items, fs=10):
    n = len(items)
    gap = 0.18
    bw = (w - gap * (n - 1)) / n
    maxh = 0
    for i, (title, body) in enumerate(items):
        bx = x + i * (bw + gap)
        body_txt = body.replace('\n', '\n')
        h = 0.38 + text_h(body_txt, bw - 0.3, fs, 1.5) + 0.22
        sh = roundbox(slide, bx, y, bw, h, BOX_BG, DARK, 0.9)
        tf = sh.text_frame
        tf.margin_right = tf.margin_left = Inches(0.12)
        tf.margin_top = tf.margin_bottom = Inches(0.06)
        p = tf.paragraphs[0]
        rtl(p)
        p.space_after = Pt(4)
        write(p, title, fs + 0.5, True, BLACK)
        for line in body.split('\n'):
            pb = tf.add_paragraph()
            rtl(pb)
            pb.space_after = Pt(2)
            write(pb, line, fs, False, BLACK)
        maxh = max(maxh, h)
    return maxh


def stat_boxes(slide, x, y, w, items):
    n = len(items)
    gap = 0.18
    bw = (w - gap * (n - 1)) / n
    h = 1.12
    for i, (val, lab, sub) in enumerate(items):
        bx = x + i * (bw + gap)
        sh = roundbox(slide, bx, y, bw, h, RGBColor(0xFF, 0xFF, 0xFF), DARK, 1.1)
        tf = sh.text_frame
        tf.margin_right = tf.margin_left = Inches(0.12)
        tf.margin_top = Inches(0.05)
        p = tf.paragraphs[0]
        rtl(p)
        p.space_after = Pt(1)
        write(p, val, 24, True, BLACK)
        p2 = tf.add_paragraph()
        rtl(p2)
        p2.space_after = Pt(1)
        write(p2, lab, 9.5, True, BLACK)
        p3 = tf.add_paragraph()
        rtl(p3)
        write(p3, sub, 8, False, GRAY)
    return h


def title_bar(slide, tab, title, page_no):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
                                 Inches(SLIDE_W), Inches(0.82))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(0x1C, 0x1C, 0x1C)
    bar.line.fill.background()
    bar.shadow.inherit = False
    tf = bar.text_frame
    tf.word_wrap = True
    tf.margin_right = tf.margin_left = Inches(0.45)
    tf.margin_top = Inches(0.10)
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    rtl(p)
    p.space_after = Pt(0)
    write(p, title, 17, True, WHITE)
    p2 = tf.add_paragraph()
    rtl(p2)
    write(p2, tab, 9.5, False, RGBColor(0xBF, 0xBF, 0xBF))
    # شماره صفحه
    textbox(slide, SLIDE_W - 1.5, SLIDE_H - 0.36, 1.1, 0.28,
            'صفحه %d از ۹' % page_no, 8.5, False, GRAY)
    textbox(slide, M, SLIDE_H - 0.36, CONTENT_W - 1.6, 0.28, FOOTER, 8, False, GRAY)


def main():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    blank = prs.slide_layouts[6]

    # ── اسلاید جلد
    s = prs.slides.add_slide(blank)
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
                              Inches(SLIDE_W), Inches(2.6))
    band.fill.solid()
    band.fill.fore_color.rgb = RGBColor(0x1C, 0x1C, 0x1C)
    band.line.fill.background()
    band.shadow.inherit = False
    tf = band.text_frame
    tf.word_wrap = True
    tf.margin_right = tf.margin_left = Inches(0.6)
    tf.margin_top = Inches(0.7)
    p = tf.paragraphs[0]
    rtl(p)
    write(p, COVER['title'], 34, True, WHITE)
    p2 = tf.add_paragraph()
    rtl(p2)
    p2.space_before = Pt(10)
    write(p2, COVER['subtitle'], 13, False, RGBColor(0xCC, 0xCC, 0xCC))
    textbox(s, M, 3.0, CONTENT_W, 0.5, COVER['tags'], 11, False, GRAY)
    for i, (t, d) in enumerate([
        ('بخش ۱', 'تداخلات دارو–دارو (صفحه ۲ و ۳)'),
        ('بخش ۲', 'تداخل دارو با غذا (صفحه ۴)'),
        ('بخش ۳', 'تداخل دارو با آزمایش (صفحه ۵)'),
        ('بخش ۴', 'ناسازگاری داروهای تزریقی (صفحه ۶)'),
        ('بخش ۵', 'گیاهان دارویی و مکمل‌ها (صفحه ۷)'),
        ('بخش ۶', 'چک‌لیست، پرچم‌های قرمز و منابع (صفحه ۸)')]):
        yy = 3.9 + i * 0.42
        textbox(s, M, yy, 1.4, 0.35, t, 12, True, BLACK)
        textbox(s, M + 1.5, yy, 6.0, 0.35, d, 11, False, GRAY)

    # ── اسلایدهای محتوا
    for pi, page in enumerate(PAGES):
        s = prs.slides.add_slide(blank)
        title_bar(s, page['tab'], page['title'], pi + 2)
        y = TOP
        for blk in page['blocks']:
            kind = blk[0]
            if kind == 'p':
                txt = blk[1]
                h = text_h(txt, CONTENT_W, 11.5, 1.4)
                textbox(s, M, y, CONTENT_W, h + 0.1, txt, 11.5, False, BLACK, spacing=1.35)
                y += h + 0.16
            elif kind == 'h':
                textbox(s, M, y, CONTENT_W, 0.34, blk[1], 14, True, BLACK)
                underline = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(M), Inches(y + 0.33),
                                               Inches(CONTENT_W), Inches(0.02))
                underline.fill.solid()
                underline.fill.fore_color.rgb = DARK
                underline.line.fill.background()
                underline.shadow.inherit = False
                y += 0.46
            elif kind == 't':
                h = add_table(s, M, y, CONTENT_W, blk[1], blk[2], blk[3])
                y += h + 0.18
            elif kind == 'pair':
                gap = 0.24
                hw = (CONTENT_W - gap) / 2
                h1 = add_table(s, M, y, hw, blk[1][0], blk[1][1], [55, 45])
                h2 = add_table(s, M + hw + gap, y, hw, blk[2][0], blk[2][1], [55, 45])
                y += max(h1, h2) + 0.18
            elif kind == 'cards':
                y += card_boxes(s, M, y, CONTENT_W, blk[1]) + 0.18
            elif kind == 'stats':
                y += stat_boxes(s, M, y, CONTENT_W, blk[1]) + 0.18
            elif kind == 'legend':
                line = '     '.join('[%s] %s' % (c, lab) for c, lab in SEVERITY)
                textbox(s, M, y, CONTENT_W, 0.3, 'راهنمای نشانگر شدت در جدول‌ها:  ' + line,
                        10, True, GRAY)
                y += 0.4
            elif kind == 'note':
                y += note_box(s, M, y, CONTENT_W, blk[1], blk[2]) + 0.16
            elif kind == 'steps':
                n = len(blk[1])
                half = (n + 1) // 2
                colw = (CONTENT_W - 0.4) / 2
                for ci, group in enumerate([blk[1][:half], blk[1][half:]]):
                    txt = '\n'.join('%d. %s' % (i + 1 + ci * half, st) for i, st in enumerate(group))
                    textbox(s, M + ci * (colw + 0.4), y, colw, text_h(txt, colw, 10.5, 1.45) + 0.1,
                            txt, 10.5, False, BLACK, spacing=1.35)
                y += max(text_h('\n'.join(blk[1][:half]), colw, 11, 1.5),
                         text_h('\n'.join(blk[1][half:]), colw, 11, 1.5)) + 0.2
            elif kind == 'refs':
                txt = '\n'.join(blk[1])
                h = text_h(txt, CONTENT_W, 9, 1.38)
                textbox(s, M, y, CONTENT_W, h + 0.1, txt, 9, False, GRAY, spacing=1.30)
                y += h + 0.06
            elif kind == 'sign':
                n = len(blk[1])
                gap = 0.2
                bw = (CONTENT_W - gap * (n - 1)) / n
                for i, lab in enumerate(blk[1]):
                    roundbox(s, M + i * (bw + gap), y, bw, 0.46, BOX_BG, DARK, 0.9)
                    textbox(s, M + i * (bw + gap) + 0.1, y + 0.09, bw - 0.2, 0.3, lab, 10, False, BLACK)
                y += 0.44

    prs.save(OUT)
    print('saved ->', OUT)

    # بررسی خروج از محدوده اسلاید
    issues = 0
    for i, s in enumerate(prs.slides):
        for sh in s.shapes:
            bottom = (sh.top or 0) + (sh.height or 0)
            right = (sh.left or 0) + (sh.width or 0)
            is_footer = (sh.top or 0) > Inches(SLIDE_H - 0.5)
            if (bottom > Inches(SLIDE_H) + 1000 or right > Inches(SLIDE_W) + 1000
                    or (bottom > Inches(SLIDE_H - 0.4) and not is_footer)):
                issues += 1
                print('  ! اسلاید %d: %.2f اینچ — %s' % (i + 1, bottom / 914400.0,
                                                      (sh.text_frame.text[:40].replace('\n', ' ')
                                                       if sh.has_text_frame else sh.shape_type)))
    print('تعداد موارد خروج از محدوده:', issues)


if __name__ == '__main__':
    main()
