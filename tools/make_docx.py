# -*- coding: utf-8 -*-
"""ساخت نسخه Word (docx) قابل ویرایش — A4 افقی، راست‌به‌چپ، فونت Arad."""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from content import COVER, PAGES, FOOTER, SEVERITY

OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                   'بروشور-تداخلات-دارویی.docx')
FONT = 'Arad'
DARK = RGBColor(0x1C, 0x1C, 0x1C)
GRAY = RGBColor(0x55, 0x55, 0x55)


def set_rtl(par):
    pPr = par._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    pPr.append(bidi)
    par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    return par


def set_font(run, size=10, bold=False, color=None, font=FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rPr.append(rf)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(attr), font)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)


def box_border(par, size=6, color='999999'):
    pPr = par._p.get_or_add_pPr()
    borders = OxmlElement('w:pBdr')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), str(size))
        el.set(qn('w:space'), '4')
        el.set(qn('w:color'), color)
        borders.append(el)
    pPr.append(borders)


def para(doc, text, size=10.5, bold=False, color=None, space_after=6, align_right=True):
    p = doc.add_paragraph()
    set_rtl(p)
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    set_font(r, size, bold, color)
    return p


def heading(doc, text, size=14):
    p = doc.add_paragraph()
    set_rtl(p)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_font(r, size, True, DARK)
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement('w:pBdr')
    bot = OxmlElement('w:bottom')
    bot.set(qn('w:val'), 'single')
    bot.set(qn('w:sz'), '8')
    bot.set(qn('w:space'), '2')
    bot.set(qn('w:color'), '444444')
    borders.append(bot)
    pPr.append(borders)
    return p


def add_table(doc, header, rows, widths=None):
    ncol = len(header)
    t = doc.add_table(rows=1, cols=ncol)
    t.style = 'Table Grid'
    t.alignment = 1
    # جهت راست‌به‌چپِ جدول
    tblPr = t._tbl.tblPr
    bidi = OxmlElement('w:bidi')
    tblPr.append(bidi)
    hdr = t.rows[0].cells
    for i, h in enumerate(header):
        hdr[i].text = ''
        p = hdr[i].paragraphs[0]
        set_rtl(p)
        r = p.add_run(h)
        set_font(r, 9.5, True, DARK)
        shade(hdr[i], 'E3E3E3')
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ''
            p = cells[i].paragraphs[0]
            set_rtl(p)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(str(val))
            bold = (i == 0 and str(val) in ('X', 'D', 'C', 'B'))
            set_font(r, 9, bold, DARK)
            if ri % 2 == 1:
                shade(cells[i], 'F5F5F5')
    if widths:
        from docx.shared import Cm as _Cm
        total = sum(widths)
        avail = 27.0  # سانتی‌متر
        for row in t.rows:
            for i, wd in enumerate(widths):
                row.cells[i].width = _Cm(avail * wd / total)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def note_box(doc, title, text):
    p = doc.add_paragraph()
    set_rtl(p)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    set_font(r, 11, True, DARK)
    box_border(p, 8, '777777')
    p2 = doc.add_paragraph()
    set_rtl(p2)
    p2.paragraph_format.space_after = Pt(10)
    p2.paragraph_format.left_indent = Cm(0.4)
    r2 = p2.add_run(text)
    set_font(r2, 10, False, None)
    box_border(p2, 8, '777777')
    return p2


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21)
    sec.left_margin = sec.right_margin = Cm(1.5)
    sec.top_margin = Cm(1.3)
    sec.bottom_margin = Cm(1.1)

    style = doc.styles['Normal']
    style.font.name = FONT
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn('w:cs'), FONT)

    # پانویس
    footer_p = sec.footer.paragraphs[0]
    set_rtl(footer_p)
    fr = footer_p.add_run(FOOTER)
    set_font(fr, 8, False, GRAY)

    # جلد
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    set_rtl(p)
    r = p.add_run(COVER['title'])
    set_font(r, 30, True, DARK)
    p = doc.add_paragraph()
    set_rtl(p)
    r = p.add_run(COVER['subtitle'])
    set_font(r, 13, False, GRAY)
    p = doc.add_paragraph()
    set_rtl(p)
    r = p.add_run(COVER['tags'])
    set_font(r, 10, False, GRAY)
    doc.add_page_break()

    for pi, page in enumerate(PAGES):
        # سربرگ صفحه
        p = doc.add_paragraph()
        set_rtl(p)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(page['tab'])
        set_font(r, 9, False, GRAY)
        p = doc.add_paragraph()
        set_rtl(p)
        p.paragraph_format.space_after = Pt(10)
        r = p.add_run(page['title'])
        set_font(r, 15, True, DARK)

        for blk in page['blocks']:
            kind = blk[0]
            if kind == 'p':
                para(doc, blk[1], 10.5, space_after=8)
            elif kind == 'h':
                heading(doc, blk[1], 13)
            elif kind == 't':
                add_table(doc, blk[1], blk[2], blk[3])
            elif kind == 'pair':
                (h1, r1), (h2, r2) = blk[1], blk[2]
                add_table(doc, h1, r1, [50, 50])
                para(doc, '', 4, space_after=2)
                add_table(doc, h2, r2, [50, 50])
            elif kind == 'cards':
                t = doc.add_table(rows=1, cols=len(blk[1]))
                t.style = 'Table Grid'
                t._tbl.tblPr.append(OxmlElement('w:bidi'))
                for i, (title, body) in enumerate(blk[1]):
                    cell = t.rows[0].cells[i]
                    cell.text = ''
                    pc = cell.paragraphs[0]
                    set_rtl(pc)
                    rr = pc.add_run(title)
                    set_font(rr, 10.5, True, DARK)
                    for line in body.split('\n'):
                        pb = cell.add_paragraph()
                        set_rtl(pb)
                        rb = pb.add_run(line)
                        set_font(rb, 9.5, False, None)
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
            elif kind == 'stats':
                t = doc.add_table(rows=2, cols=len(blk[1]))
                t.style = 'Table Grid'
                t._tbl.tblPr.append(OxmlElement('w:bidi'))
                for i, (v, lab, sub) in enumerate(blk[1]):
                    c0, c1 = t.rows[0].cells[i], t.rows[1].cells[i]
                    c0.text = ''
                    p0 = c0.paragraphs[0]
                    set_rtl(p0)
                    r0 = p0.add_run(v)
                    set_font(r0, 22, True, DARK)
                    c1.text = ''
                    p1 = c1.paragraphs[0]
                    set_rtl(p1)
                    r1 = p1.add_run(lab)
                    set_font(r1, 9.5, True, None)
                    p1b = c1.add_paragraph()
                    set_rtl(p1b)
                    r1b = p1b.add_run(sub)
                    set_font(r1b, 8.5, False, GRAY)
                doc.add_paragraph().paragraph_format.space_after = Pt(4)
            elif kind == 'legend':
                line = '   |   '.join('%s = %s' % (c, lab) for c, lab in SEVERITY)
                para(doc, 'راهنمای نشانگر شدت:  ' + line, 9.5, False, GRAY, space_after=8)
            elif kind == 'note':
                note_box(doc, blk[1], blk[2])
            elif kind == 'steps':
                for i, st in enumerate(blk[1], 1):
                    para(doc, '%d. %s' % (i, st), 10.5, space_after=3)
            elif kind == 'refs':
                for r_ in blk[1]:
                    para(doc, r_, 9, False, GRAY, space_after=3)
            elif kind == 'sign':
                t = doc.add_table(rows=1, cols=len(blk[1]))
                t.style = 'Table Grid'
                t._tbl.tblPr.append(OxmlElement('w:bidi'))
                for i, lab in enumerate(blk[1]):
                    cell = t.rows[0].cells[i]
                    cell.text = ''
                    pc = cell.paragraphs[0]
                    set_rtl(pc)
                    rr = pc.add_run(lab)
                    set_font(rr, 10, False, DARK)
                doc.add_paragraph().paragraph_format.space_after = Pt(4)
        if pi < len(PAGES) - 1:
            doc.add_page_break()

    doc.save(OUT)
    print('saved ->', OUT)


if __name__ == '__main__':
    main()
